"""
Client Profile Generator Service.
Transforms candidate CVs and JDs into structured, anonymized client-ready profiles
and generates downloadable 16:9 widescreen PowerPoint presentation decks (.pptx)
matching Protiviti's exact single-slide executive candidate profile format.
"""

import io
import re
import json
from openai import AsyncOpenAI

from services.jd_generator import call_llm_with_fallback

CLIENT_PROFILE_SYSTEM_PROMPT = """You are an experienced Talent Transformation and Recruitment Consultant preparing ATS-friendly candidate profiles for client submissions.

INPUTS:
1. Candidate CV
2. Job Description (JD)

OBJECTIVE:
Create a concise, fact-based, client-ready submission profile that aligns the candidate's experience to the target JD using only verifiable facts from the CV.

CORE ANONYMIZATION & FACT RULES:
- Remove employer names and replace them with industry descriptions (e.g., Global Bank, Technology Services Company, Global Asset Management Firm).
- Do NOT fabricate achievements, tools, certifications, or business outcomes.
- Use terminology from the JD whenever supported by the CV.

EXPERIENCE CALCULATION RULES (CAREER GAPS):
- Calculate total active working experience from employment start and end dates (Month/Year).
- SUBTRACT CAREER GAPS: If there are gaps/sabbaticals between roles, subtract the gap durations.
- State the exact net active years in both `experience_years` and the opening sentence of `professional_summary` (e.g. "<Candidate Name> has X years of experience...").

TASK INSTRUCTIONS:
1. PROFESSIONAL SUMMARY: 120–150 words highlighting functional/domain expertise, key tools, and business impact. Start with "<Candidate Name> has X years of experience...".
2. INDUSTRY EXPERTISE: 2–3 relevant industry domains supported by the CV (e.g., Capital Markets, Digital Transformation).
3. MAJOR PROJECT EXPERIENCE:
   - Create exactly 4 major projects (or up to 4 for shorter CVs).
   - Provide 3–4 high-impact bullets per project (target 18–22 words per bullet).
   - Total Major Projects text MUST NOT exceed 2,600 characters including spaces.
   - Preserve only the strongest, non-repetitive experience and keep business meaning intact.
4. AREAS OF EXPERTISE: Exactly 20 ATS keywords (skills, tools, methodologies) present in the CV.
5. EDUCATION & GRADE: Top 3 degrees or major certifications ordered highest to lowest. Extract grade/band if stated.

OUTPUT FORMAT INSTRUCTIONS:
- Return ONLY a raw JSON object starting with { and ending with }.
- Do NOT include markdown code fences like ```json or ```.
- Ensure all quotes within text strings are properly escaped (e.g. use \\" if needed).

JSON STRUCTURE:
{
  "candidate_name": "Candidate Name",
  "grade_band": "GCB 4",
  "experience_years": "8+ years",
  "professional_summary": "Candidate Name has 8 years of experience in...",
  "industry_expertise": [
    "Regulatory Reporting",
    "Capital Markets"
  ],
  "education": [
    "Bachelor of Engineering – Computer Science"
  ],
  "areas_of_expertise": [
    "Skill 1", "Skill 2", "Skill 3", "Skill 4", "Skill 5",
    "Skill 6", "Skill 7", "Skill 8", "Skill 9", "Skill 10",
    "Skill 11", "Skill 12", "Skill 13", "Skill 14", "Skill 15",
    "Skill 16", "Skill 17", "Skill 18", "Skill 19", "Skill 20"
  ],
  "major_projects": [
    {
      "industry_desc": "Global Investment Bank",
      "project_theme": "Regulatory Reporting Transformation",
      "role": "BA & Test Analyst",
      "bullets": [
        "Bullet point 1 detailing project contributions...",
        "Bullet point 2 detailing tools and impact...",
        "Bullet point 3 detailing validation..."
      ]
    }
  ]
}"""


def extract_clean_json(text: str) -> dict:
    """Safely extracts and parses JSON even if wrapped in markdown, containing
    unescaped characters, single quotes instead of double quotes, or truncated mid-output by token limits."""
    import ast
    if not text or not text.strip():
        raise ValueError("Empty response received from LLM model. The model may be overloaded — please try again.")
    
    clean_text = text.strip()
    
    # Stage 1: Strip markdown backticks ```json ... ```
    if "```" in clean_text:
        clean_text = re.sub(r"^```(?:json)?\s*", "", clean_text, flags=re.MULTILINE)
        clean_text = re.sub(r"\s*```\s*$", "", clean_text, flags=re.MULTILINE)
        clean_text = clean_text.strip()
    
    # Stage 2: Extract outer JSON braces if extra commentary exists
    json_match = re.search(r"(\{.*)", clean_text, re.DOTALL)
    if json_match:
        clean_text = json_match.group(1).strip()
    
    # Stage 3: Try standard json.loads parsing directly
    try:
        return json.loads(clean_text)
    except json.JSONDecodeError:
        pass
    
    # Stage 4: Try ast.literal_eval parsing (safely handles single quotes and double quotes together)
    try:
        parsed_dict = ast.literal_eval(clean_text)
        if isinstance(parsed_dict, dict):
            return parsed_dict
    except Exception:
        pass

    # Stage 5: Repair unescaped newlines inside string values
    try:
        repaired = re.sub(r'(?<!\\)\n', r'\\n', clean_text)
        return json.loads(repaired)
    except json.JSONDecodeError:
        pass
    
    # Stage 6: Auto-close truncated JSON (model ran out of tokens mid-output)
    try:
        truncated = clean_text.rstrip()
        # Remove trailing comma if present
        truncated = re.sub(r',\s*$', '', truncated)
        # Count unclosed brackets and braces
        open_braces = truncated.count('{') - truncated.count('}')
        open_brackets = truncated.count('[') - truncated.count(']')
        # Check if we're inside an unclosed string — close it
        quote_count = truncated.count('"') - truncated.count('\\"')
        if quote_count % 2 != 0:
            truncated += '"'
        # Close arrays then objects
        truncated += ']' * max(0, open_brackets)
        truncated += '}' * max(0, open_braces)
        repaired = re.sub(r'(?<!\\)\n', r'\\n', truncated)
        return json.loads(repaired)
    except json.JSONDecodeError:
        pass

    # Stage 7: Try combining ast.literal_eval + auto-close repair
    try:
        truncated = clean_text.rstrip()
        truncated = re.sub(r',\s*$', '', truncated)
        open_braces = truncated.count('{') - truncated.count('}')
        open_brackets = truncated.count('[') - truncated.count(']')
        
        # Check single quotes vs double quotes balance
        sq_count = truncated.count("'") - truncated.count("\\'")
        if sq_count % 2 != 0:
            truncated += "'"
        else:
            dq_count = truncated.count('"') - truncated.count('\\"')
            if dq_count % 2 != 0:
                truncated += '"'
                
        truncated += ']' * max(0, open_brackets)
        truncated += '}' * max(0, open_braces)
        
        parsed_dict = ast.literal_eval(truncated)
        if isinstance(parsed_dict, dict):
            return parsed_dict
    except Exception as e:
        raise ValueError(f"Failed to parse LLM response as JSON after all repair attempts: {str(e)[:200]}")


async def generate_client_profile(client: AsyncOpenAI, resume_text: str, jd_text: str) -> dict:
    """Calls LLM to generate structured Client Submission Profile JSON."""
    # Truncate prompt context if CV or JD is excessively long to protect TPM limits
    safe_resume = resume_text[:12000] if resume_text else ""
    safe_jd = jd_text[:5000] if jd_text else ""
    jd_header = f"--- JOB DESCRIPTION (JD) ---\n{safe_jd}" if safe_jd.strip() else "--- JOB DESCRIPTION (JD) ---\nNo specific Job Description provided. Extract and format the profile purely from the Candidate CV."

    user_msg = (
        f"{jd_header}\n\n"
        f"--- CANDIDATE CV ---\n{safe_resume}"
    )

    try:
        result_text = await call_llm_with_fallback(
            client=client,
            messages=[
                {"role": "system", "content": CLIENT_PROFILE_SYSTEM_PROMPT},
                {"role": "user", "content": user_msg},
            ],
            temperature=0.2,
            max_tokens=3500,
        )
        data = extract_clean_json(result_text)

        # Fallback validation for missing keys
        if "areas_of_expertise" not in data: data["areas_of_expertise"] = []
        if "industry_expertise" not in data: data["industry_expertise"] = []
        if "education" not in data: data["education"] = []
        if "major_projects" not in data: data["major_projects"] = []
        if "grade_band" not in data: data["grade_band"] = ""
        
        return data
    except Exception as e:
        return {
            "candidate_name": "Candidate Profile",
            "grade_band": "",
            "experience_years": "N/A",
            "professional_summary": f"Error generating client profile: {str(e)}",
            "industry_expertise": [],
            "education": [],
            "major_projects": [],
            "areas_of_expertise": [],
        }


def build_client_profile_html(profile_data: dict) -> str:
    """
    Renders an HTML slide with exact 16:9 widescreen layout (13.333 in x 7.5 in equivalent),
    featuring exact Protiviti color tokens, typography, dark blue left banner, cyan headers,
    and responsive flex columns.
    """
    def clean_text(text: str) -> str:
        if not isinstance(text, str): return ""
        return text.replace("**", "").replace("*", "").strip()

    cand_name = clean_text(profile_data.get("candidate_name", "Candidate Profile"))
    grade_band = clean_text(profile_data.get("grade_band", "GCB 4"))
    summary = clean_text(profile_data.get("professional_summary", ""))
    industries = [clean_text(i) for i in profile_data.get("industry_expertise", [])]
    education = [clean_text(e) for e in profile_data.get("education", [])]
    keywords = [clean_text(k) for k in profile_data.get("areas_of_expertise", [])]
    projects = profile_data.get("major_projects", [])

    name_parts = cand_name.split(" ", 1)
    first_name = name_parts[0]
    last_name = name_parts[1] if len(name_parts) > 1 else ""

    summary_clean = summary.strip()
    if ". " in summary_clean:
        first_sent, rest_summary = summary_clean.split(". ", 1)
        formatted_summary = f"<strong>{first_sent}.</strong> {rest_summary}"
    else:
        formatted_summary = f"<strong>{summary_clean}</strong>"

    kw_lis = "".join([f"<li style='margin-bottom:5px; line-height:1.4;'>{kw}</li>" for kw in keywords[:20]])
    ind_lis = "".join([f"<li style='margin-bottom:5px; line-height:1.4;'>{ind}</li>" for ind in industries[:3]])
    ed_lis = "".join([f"<li style='margin-bottom:5px; line-height:1.4;'>{ed}</li>" for ed in (education if education else ["Bachelor Degree in Engineering / Related Field"])[:3]])

    proj_blocks = ""
    for proj in projects[:4]:
        ind_desc = clean_text(proj.get("industry_desc", "Global Organization"))
        theme = clean_text(proj.get("project_theme", "Strategic Project"))
        role = clean_text(proj.get("role", "Consultant"))
        bullets = [clean_text(b) for b in proj.get("bullets", [])]
        bullets_lis = "".join([f"<li style='margin-bottom:2px;'>{b}</li>" for b in bullets[:5]])
        
        proj_blocks += f"""
        <div style="margin-bottom: 8px;">
            <div style="font-weight: bold; color: #1e293b; font-size: 13px; margin-bottom: 2px;">{ind_desc} – {theme} – {role}</div>
            <ul style="margin: 0; padding-left: 16px; color: #334155; font-size: 12px; line-height: 1.35;">
                {bullets_lis}
            </ul>
        </div>
        """

    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  @import url('https://fonts.googleapis.com/css2?family=Barlow:wght@400;600;700;800&display=swap');
  @page {
    size: 13.333in 7.5in;
    margin: 0;
  }
  * {
    box-sizing: border-box;
    font-family: 'Barlow', Arial, sans-serif;
  }
  body {
    margin: 0;
    padding: 0;
    width: 13.333in;
    height: 7.5in;
    background: #ffffff;
    display: flex;
    overflow: hidden;
  }
  .left-banner {
    width: 3.2in;
    height: 7.5in;
    background: #0B3B60;
    color: #ffffff;
    padding: 2.2in 0.4in 0.4in 0.4in;
    display: flex;
    flex-direction: column;
  }
  .cand-first-name {
    font-size: 32px;
    font-weight: bold;
    line-height: 1.1;
    text-transform: uppercase;
  }
  .cand-last-name {
    font-size: 32px;
    font-weight: bold;
    line-height: 1.1;
    margin-bottom: 36px;
    text-transform: uppercase;
  }
  .grade-band {
    font-size: 16px;
    font-weight: bold;
    color: #ffffff;
  }
  .main-content {
    flex: 1;
    height: 7.5in;
    padding: 0.3in 0.4in;
    display: flex;
    gap: 0.3in;
    position: relative;
  }
  .col1 {
    width: 2.7in;
    display: flex;
    flex-direction: column;
  }
  .col2 {
    flex: 1;
    display: flex;
    flex-direction: column;
  }
  .section-header {
    font-size: 13px;
    font-weight: bold;
    color: #0096C8;
    text-transform: uppercase;
    margin-bottom: 4px;
    letter-spacing: 0.5px;
  }
  ul.bullet-list {
    margin: 0 0 12px 0;
    padding-left: 14px;
    color: #1e293b;
    font-size: 12px; /* 9pt equivalent */
    line-height: 1.35;
  }
  .summary-text {
    font-size: 12px; /* 9pt equivalent */
    color: #1e293b;
    line-height: 1.4;
    margin-bottom: 12px;
  }
  .brand-logo {{
    position: absolute;
    bottom: 0.3in;
    right: 0.4in;
    font-size: 22px;
    font-weight: bold;
    color: #0B3B60;
    letter-spacing: -0.5px;
  }}
</style>
</head>
<body>
  <div class="left-banner">
    <div class="cand-first-name">{first_name}</div>
    <div class="cand-last-name">{last_name}</div>
    <div class="grade-band">{grade_band}</div>
  </div>
  <div class="main-content">
    <div class="col1">
      <div class="section-header">Areas of Expertise</div>
      <ul class="bullet-list">
        {kw_lis}
      </ul>
      <div class="section-header" style="margin-top: 8px;">Industry Expertise</div>
      <ul class="bullet-list">
        {ind_lis}
      </ul>
      <div class="section-header" style="margin-top: 8px;">Education</div>
      <ul class="bullet-list">
        {ed_lis}
      </ul>
    </div>
    <div class="col2">
      <div class="section-header">Professional Experience</div>
      <div class="summary-text">{formatted_summary}</div>
      <div class="section-header" style="margin-top: 4px;">Major Projects</div>
      {proj_blocks}
    </div>
    <div class="brand-logo">protiviti</div>
  </div>
</body>
</html>
"""
    return html


def build_client_pptx_deck(profile_data: dict) -> bytes:
    """
    Generates a single 16:9 widescreen PowerPoint slide matching Protiviti's exact design format.
    Uses templates/client_profile_template.pptx if present to preserve exact layout and typography,
    otherwise falls back to generating the slide programmatically.
    """
    try:
        import os
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        raise RuntimeError("The 'python-pptx' library is not installed in your Python environment. Please run: pip install python-pptx")

    template_path = os.path.join(os.path.dirname(__file__), "..", "templates", "client_profile_template.pptx")
    if not os.path.exists(template_path):
        template_path = os.path.join(os.path.dirname(__file__), "..", "templates", "client_profile_template.pptx.pptx")

    def clean_text(text: str) -> str:
        if not isinstance(text, str): return ""
        return text.replace("**", "").replace("*", "").strip()

    cand_name = clean_text(profile_data.get("candidate_name", "Candidate Profile"))
    grade_band = clean_text(profile_data.get("grade_band", ""))
    if grade_band.lower() in ["not mentioned", "n/a", "none"]:
        grade_band = ""
    summary = clean_text(profile_data.get("professional_summary", ""))
    industries = [clean_text(i) for i in profile_data.get("industry_expertise", [])]
    education = [clean_text(e) for e in profile_data.get("education", [])]
    keywords = [clean_text(k) for k in profile_data.get("areas_of_expertise", [])]
    projects = profile_data.get("major_projects", [])

    if os.path.exists(template_path):
        prs = Presentation(template_path)
        slide = prs.slides[0]

        def set_p_text(p, text, bold=None, size_pt=9.0, color_rgb=None):
            p.text = text
            p.font.name = "Barlow"
            p.font.size = Pt(size_pt)
            if bold is not None:
                p.font.bold = bold
            if color_rgb is not None:
                p.font.color.rgb = color_rgb

        HEADER_CYAN = RGBColor(0, 150, 200)
        TEXT_DARK = RGBColor(30, 30, 30)
        WHITE = RGBColor(255, 255, 255)

        shapes = [s for s in slide.shapes if s.has_text_frame]
        shapes_sorted = sorted(shapes, key=lambda s: s.left)

        # Flag to prevent double-populating the same column across multiple shapes
        populated_col0 = False
        populated_col1 = False
        populated_col2 = False

        for shape in shapes_sorted:
            tf = shape.text_frame
            full_text = tf.text.upper()

            # Column 0: Left Banner (Candidate Name & Grade)
            if not populated_col0 and ("{{CANDIDATE_NAME}}" in full_text or "CANDIDATE" in full_text or shape.left < Inches(3.3)):
                populated_col0 = True
                tf.clear()
                name_parts = cand_name.split(" ", 1)
                p_name1 = tf.paragraphs[0]
                set_p_text(p_name1, name_parts[0], bold=True, size_pt=32, color_rgb=WHITE)
                if len(name_parts) > 1:
                    p_name2 = tf.add_paragraph()
                    set_p_text(p_name2, name_parts[1], bold=True, size_pt=32, color_rgb=WHITE)
                if grade_band:
                    p_g = tf.add_paragraph()
                    set_p_text(p_g, grade_band, bold=True, size_pt=16, color_rgb=WHITE)
                    p_g.space_before = Pt(24)

            # Column 1: Areas of Expertise, Industry, Education
            elif not populated_col1 and ("{{AREAS_OF_EXPERTISE}}" in full_text or "AREAS OF EXPERTISE" in full_text or (shape.left >= Inches(3.3) and shape.left < Inches(6.0))):
                populated_col1 = True
                tf.clear()
                p_hdr = tf.paragraphs[0]
                set_p_text(p_hdr, "AREAS OF EXPERTISE", bold=True, size_pt=9.5, color_rgb=HEADER_CYAN)
                p_hdr.space_after = Pt(3)
                for kw in keywords[:20]:
                    p_k = tf.add_paragraph()
                    set_p_text(p_k, f"• {kw}", size_pt=9.0, color_rgb=TEXT_DARK)
                    p_k.space_after = Pt(3.5)
                if industries:
                    p_ie = tf.add_paragraph()
                    set_p_text(p_ie, "INDUSTRY EXPERTISE", bold=True, size_pt=9.5, color_rgb=HEADER_CYAN)
                    p_ie.space_before = Pt(10)
                    p_ie.space_after = Pt(3)
                    for ind in industries[:3]:
                        p_i = tf.add_paragraph()
                        set_p_text(p_i, f"• {ind}", size_pt=9.0, color_rgb=TEXT_DARK)
                        p_i.space_after = Pt(3.5)
                if education:
                    p_ed = tf.add_paragraph()
                    set_p_text(p_ed, "EDUCATION", bold=True, size_pt=9.5, color_rgb=HEADER_CYAN)
                    p_ed.space_before = Pt(10)
                    p_ed.space_after = Pt(3)
                    for ed in education[:3]:
                        p_e = tf.add_paragraph()
                        set_p_text(p_e, f"• {ed}", size_pt=9.0, color_rgb=TEXT_DARK)
                        p_e.space_after = Pt(3.5)

            # Column 2: Professional Experience & Major Projects
            elif not populated_col2 and ("{{PROFESSIONAL_EXPERIENCE}}" in full_text or "PROFESSIONAL EXPERIENCE" in full_text or shape.left >= Inches(6.0)):
                populated_col2 = True
                tf.clear()
                p_pe = tf.paragraphs[0]
                set_p_text(p_pe, "PROFESSIONAL EXPERIENCE", bold=True, size_pt=9.5, color_rgb=HEADER_CYAN)
                p_pe.space_after = Pt(3)
                
                p_sum = tf.add_paragraph()
                # Split summary into first sentence and remaining body to bold first line
                summary_clean = summary.strip()
                if ". " in summary_clean:
                    first_sent, rest_summary = summary_clean.split(". ", 1)
                    run_first = p_sum.add_run()
                    run_first.text = first_sent + ". "
                    run_first.font.name = "Barlow"
                    run_first.font.size = Pt(9.0)
                    run_first.font.bold = True
                    run_first.font.color.rgb = TEXT_DARK

                    run_rest = p_sum.add_run()
                    run_rest.text = rest_summary
                    run_rest.font.name = "Barlow"
                    run_rest.font.size = Pt(9.0)
                    run_rest.font.bold = False
                    run_rest.font.color.rgb = TEXT_DARK
                else:
                    set_p_text(p_sum, summary, bold=True, size_pt=9.0, color_rgb=TEXT_DARK)
                p_sum.space_after = Pt(6)

                p_mp = tf.add_paragraph()
                set_p_text(p_mp, "MAJOR PROJECTS", bold=True, size_pt=9.5, color_rgb=HEADER_CYAN)
                p_mp.space_before = Pt(6)
                p_mp.space_after = Pt(4)

                for proj in projects[:4]:
                    ind_desc = clean_text(proj.get("industry_desc", "Global Organization"))
                    theme = clean_text(proj.get("project_theme", "Strategic Project"))
                    role = clean_text(proj.get("role", "Consultant"))
                    bullets = [clean_text(b) for b in proj.get("bullets", [])]

                    p_title = tf.add_paragraph()
                    set_p_text(p_title, f"{ind_desc} – {theme} – {role}", bold=True, size_pt=9.0, color_rgb=TEXT_DARK)
                    p_title.space_before = Pt(4)
                    p_title.space_after = Pt(1)

                    for b in bullets[:5]:
                        p_b = tf.add_paragraph()
                        set_p_text(p_b, f"• {b}", size_pt=9.0, color_rgb=TEXT_DARK)
                        p_b.space_after = Pt(1)

        out_stream = io.BytesIO()
        prs.save(out_stream)
        out_stream.seek(0)
        return out_stream.getvalue()

    # Fallback to standard generation if template file is not provided yet
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    BANNER_BLUE = RGBColor(11, 59, 96)
    HEADER_CYAN = RGBColor(0, 150, 200)
    TEXT_DARK = RGBColor(30, 30, 30)
    TEXT_MUTED = RGBColor(90, 90, 90)
    WHITE = RGBColor(255, 255, 255)
    FONT_NAME = "Barlow"

    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(3.2), Inches(7.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = BANNER_BLUE
    banner.line.fill.background()

    tf_name = slide.shapes.add_textbox(Inches(0.4), Inches(2.2), Inches(2.6), Inches(4.5)).text_frame
    tf_name.word_wrap = True
    tf_name.margin_left = Inches(0)
    tf_name.margin_top = Inches(0)
    
    name_parts = cand_name.split(" ", 1)
    p_name1 = tf_name.paragraphs[0]
    p_name1.text = name_parts[0]
    p_name1.font.size = Pt(32)
    p_name1.font.bold = True
    p_name1.font.color.rgb = WHITE
    p_name1.font.name = FONT_NAME

    if len(name_parts) > 1:
        p_name2 = tf_name.add_paragraph()
        p_name2.text = name_parts[1]
        p_name2.font.size = Pt(32)
        p_name2.font.bold = True
        p_name2.font.color.rgb = WHITE
        p_name2.font.name = FONT_NAME
        p_name2.space_before = Pt(0)

    if grade_band:
        p_grade = tf_name.add_paragraph()
        p_grade.text = grade_band
        p_grade.font.size = Pt(16)
        p_grade.font.bold = True
        p_grade.font.color.rgb = WHITE
        p_grade.font.name = FONT_NAME
        p_grade.space_before = Pt(36)

    col1_left = Inches(3.45)
    col1_width = Inches(2.7)

    tf_col1 = slide.shapes.add_textbox(col1_left, Inches(0.3), col1_width, Inches(6.9)).text_frame
    tf_col1.word_wrap = True
    tf_col1.margin_left = Inches(0)
    tf_col1.margin_top = Inches(0)

    p_aoe = tf_col1.paragraphs[0]
    p_aoe.text = "AREAS OF EXPERTISE"
    p_aoe.font.size = Pt(9.5)
    p_aoe.font.bold = True
    p_aoe.font.color.rgb = HEADER_CYAN
    p_aoe.font.name = FONT_NAME
    p_aoe.space_after = Pt(3)

    for kw in keywords[:20]:
        p_k = tf_col1.add_paragraph()
        p_k.text = f"• {kw}"
        p_k.font.size = Pt(9.0)
        p_k.font.color.rgb = TEXT_DARK
        p_k.font.name = FONT_NAME
        p_k.space_before = Pt(0)
        p_k.space_after = Pt(3.5)

    p_ie = tf_col1.add_paragraph()
    p_ie.text = "INDUSTRY EXPERTISE"
    p_ie.font.size = Pt(9.5)
    p_ie.font.bold = True
    p_ie.font.color.rgb = HEADER_CYAN
    p_ie.font.name = FONT_NAME
    p_ie.space_before = Pt(12)
    p_ie.space_after = Pt(3)

    for ind in industries[:3]:
        p_i = tf_col1.add_paragraph()
        p_i.text = f"• {ind}"
        p_i.font.size = Pt(9.0)
        p_i.font.color.rgb = TEXT_DARK
        p_i.font.name = FONT_NAME
        p_i.space_before = Pt(0)
        p_i.space_after = Pt(3.5)

    p_ed = tf_col1.add_paragraph()
    p_ed.text = "EDUCATION"
    p_ed.font.size = Pt(9.5)
    p_ed.font.bold = True
    p_ed.font.color.rgb = HEADER_CYAN
    p_ed.font.name = FONT_NAME
    p_ed.space_before = Pt(12)
    p_ed.space_after = Pt(3)

    ed_list = education if education else ["Bachelor Degree in Engineering / Related Field"]
    for ed in ed_list[:3]:
        p_e = tf_col1.add_paragraph()
        p_e.text = f"• {ed}"
        p_e.font.size = Pt(9.0)
        p_e.font.color.rgb = TEXT_DARK
        p_e.font.name = FONT_NAME
        p_e.space_before = Pt(0)
        p_e.space_after = Pt(3.5)

    col2_left = Inches(6.35)
    col2_width = Inches(6.6)

    tf_col2 = slide.shapes.add_textbox(col2_left, Inches(0.3), col2_width, Inches(6.6)).text_frame
    tf_col2.word_wrap = True
    tf_col2.margin_left = Inches(0)
    tf_col2.margin_top = Inches(0)

    p_pe = tf_col2.paragraphs[0]
    p_pe.text = "PROFESSIONAL EXPERIENCE"
    p_pe.font.size = Pt(9.5)
    p_pe.font.bold = True
    p_pe.font.color.rgb = HEADER_CYAN
    p_pe.font.name = FONT_NAME
    p_pe.space_after = Pt(3)

    p_sum = tf_col2.add_paragraph()
    summary_clean = summary.strip()
    if ". " in summary_clean:
        first_sent, rest_summary = summary_clean.split(". ", 1)
        r_first = p_sum.add_run()
        r_first.text = first_sent + ". "
        r_first.font.size = Pt(9.0)
        r_first.font.bold = True
        r_first.font.color.rgb = TEXT_DARK
        r_first.font.name = FONT_NAME

        r_rest = p_sum.add_run()
        r_rest.text = rest_summary
        r_rest.font.size = Pt(9.0)
        r_rest.font.bold = False
        r_rest.font.color.rgb = TEXT_DARK
        r_rest.font.name = FONT_NAME
    else:
        p_sum.text = summary
        p_sum.font.size = Pt(9.0)
        p_sum.font.bold = True
        p_sum.font.color.rgb = TEXT_DARK
        p_sum.font.name = FONT_NAME
    p_sum.space_before = Pt(0)
    p_sum.space_after = Pt(8)

    p_mp = tf_col2.add_paragraph()
    p_mp.text = "MAJOR PROJECTS"
    p_mp.font.size = Pt(9.5)
    p_mp.font.bold = True
    p_mp.font.color.rgb = HEADER_CYAN
    p_mp.font.name = FONT_NAME
    p_mp.space_before = Pt(4)
    p_mp.space_after = Pt(4)

    for proj in projects[:4]:
        ind_desc = clean_text(proj.get("industry_desc", "Global Organization"))
        theme = clean_text(proj.get("project_theme", "Strategic Project"))
        role = clean_text(proj.get("role", "Consultant"))
        bullets = [clean_text(b) for b in proj.get("bullets", [])]

        p_title = tf_col2.add_paragraph()
        p_title.text = f"{ind_desc} – {theme} – {role}"
        p_title.font.size = Pt(9.0)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.font.name = FONT_NAME
        p_title.space_before = Pt(4)
        p_title.space_after = Pt(1)

        for b in bullets[:5]:
            p_b = tf_col2.add_paragraph()
            p_b.text = f"• {b}"
            p_b.font.size = Pt(9.0)
            p_b.font.color.rgb = TEXT_DARK
            p_b.font.name = FONT_NAME
            p_b.space_before = Pt(0)
            p_b.space_after = Pt(1)

    tf_logo = slide.shapes.add_textbox(Inches(11.2), Inches(6.8), Inches(1.8), Inches(0.5)).text_frame
    p_logo = tf_logo.paragraphs[0]
    p_logo.text = "protiviti"
    p_logo.font.size = Pt(18)
    p_logo.font.bold = True
    p_logo.font.color.rgb = BANNER_BLUE
    p_logo.font.name = FONT_NAME
    p_logo.alignment = PP_ALIGN.RIGHT

    out_stream = io.BytesIO()
    prs.save(out_stream)
    out_stream.seek(0)
    return out_stream.getvalue()
